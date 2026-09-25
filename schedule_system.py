import asyncio
import os
import re
import json
import urllib.parse
import urllib.request
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
from typing import Optional

from aiogram import F, Router
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


schedule_router = Router()

_admin_ids: list[int] = []
_excel_lock = asyncio.Lock()

_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SCHEDULE_FILE = os.path.join(_BASE_DIR, "work_schedules.xlsx")
DAYOFF_FILE = os.path.join(_BASE_DIR, "day_off_requests.xlsx")
TIMEZONE_FILE = os.path.join(_BASE_DIR, "user_timezones.json")

WEEKDAYS = [
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
    "Sunday",
]

WEEKDAY_RU = {
    "Monday": "Понедельник",
    "Tuesday": "Вторник",
    "Wednesday": "Среда",
    "Thursday": "Четверг",
    "Friday": "Пятница",
    "Saturday": "Суббота",
    "Sunday": "Воскресенье",
}

WEEKDAY_SHORT_RU = {
    "Monday": "Пн",
    "Tuesday": "Вт",
    "Wednesday": "Ср",
    "Thursday": "Чт",
    "Friday": "Пт",
    "Saturday": "Сб",
    "Sunday": "Вс",
}

STATUS_RU = {
    "Working": "Работа",
    "Day Off": "Выходной",
    "Pending": "На рассмотрении",
    "Approved": "Одобрено",
    "Rejected": "Отклонено",
}

DAY_ALIASES = {
    "monday": "Monday",
    "mon": "Monday",
    "понедельник": "Monday",
    "пн": "Monday",
    "tuesday": "Tuesday",
    "tue": "Tuesday",
    "tues": "Tuesday",
    "вторник": "Tuesday",
    "вт": "Tuesday",
    "wednesday": "Wednesday",
    "wed": "Wednesday",
    "среда": "Wednesday",
    "ср": "Wednesday",
    "thursday": "Thursday",
    "thu": "Thursday",
    "thur": "Thursday",
    "thurs": "Thursday",
    "четверг": "Thursday",
    "чт": "Thursday",
    "friday": "Friday",
    "fri": "Friday",
    "пятница": "Friday",
    "пт": "Friday",
    "saturday": "Saturday",
    "sat": "Saturday",
    "суббота": "Saturday",
    "сб": "Saturday",
    "sunday": "Sunday",
    "sun": "Sunday",
    "воскресенье": "Sunday",
    "вс": "Sunday",
}

OFF_ALIASES = {
    "day off",
    "day-off",
    "off",
    "выходной",
    "вых",
    "holiday",
    "rest",
}

SCHEDULE_HEADERS = [
    "Telegram User ID",
    "Username",
    "Display Name",
    "Day of Week",
    "Start Time",
    "End Time",
    "Status",
    "Last Updated",
]

DAYOFF_HEADERS = [
    "Request ID",
    "Telegram User ID",
    "Username",
    "Display Name",
    "Requested Date",
    "Reason",
    "Status",
    "Last Updated",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")
APPROVED_FILL = PatternFill("solid", fgColor="C6EFCE")
REJECTED_FILL = PatternFill("solid", fgColor="FFC7CE")


class ScheduleStates(StatesGroup):
    waiting_city = State()
    waiting_schedule_text = State()
    waiting_day_time = State()
    waiting_dayoff_date = State()
    waiting_dayoff_reason = State()


def _load_timezones() -> dict:
    try:
        with open(TIMEZONE_FILE, "r", encoding="utf-8") as f: return json.load(f)
    except (OSError, json.JSONDecodeError): return {}

def _save_timezones(data: dict):
    with open(TIMEZONE_FILE, "w", encoding="utf-8") as f: json.dump(data, f, ensure_ascii=False, indent=2)

def user_tz_info(user_id) -> dict:
    return _load_timezones().get(str(user_id), {})

def find_city_timezone(city: str) -> dict | None:
    try:
        q=urllib.parse.urlencode({"name":city,"count":1,"language":"ru","format":"json"})
        req=urllib.request.Request("https://geocoding-api.open-meteo.com/v1/search?"+q, headers={"User-Agent":"ShinobiScheduleBot/1.0"})
        with urllib.request.urlopen(req, timeout=8) as r: data=json.load(r)
        x=(data.get("results") or [None])[0]
        if not x or not x.get("timezone"): return None
        return {"city":x.get("name",city),"country":x.get("country",""),"timezone":x["timezone"]}
    except Exception: return None

def _convert_hhmm(value: str, from_tz: str, to_tz: str="Europe/Warsaw") -> str:
    if not value: return value
    try:
        h,m=map(int,value.split(":")); now=datetime.now(ZoneInfo(from_tz)); dt=now.replace(hour=h,minute=m,second=0,microsecond=0)
        return dt.astimezone(ZoneInfo(to_tz)).strftime("%H:%M")
    except Exception: return value

def format_schedule_poland(draft: dict, user_id) -> str:
    info=user_tz_info(user_id); tz=info.get("timezone")
    if not tz: return format_schedule(draft)
    converted={}
    for day,item in draft.items():
        x=dict(item)
        if x.get("status") != "Day Off":
            x["start"]=_convert_hhmm(x.get("start",""),tz); x["end"]=_convert_hhmm(x.get("end",""),tz)
        converted[day]=x
    return format_schedule(converted)

def init_schedule_system(admin_ids: list[int]):
    global _admin_ids
    _admin_ids = list(admin_ids)


def is_schedule_admin(user_id: int) -> bool:
    return user_id in _admin_ids


def display_name(user) -> str:
    return user.full_name or "Неизвестно"


def day_ru(day: str) -> str:
    return WEEKDAY_RU.get(day, day)


def status_ru(status: str) -> str:
    return STATUS_RU.get(status, status)


def normalize_day(day: str) -> str:
    if day in WEEKDAYS:
        return day
    return DAY_ALIASES.get(str(day).strip().lower(), "")


def normalize_status(status: str) -> str:
    text = str(status or "").strip().lower()
    if text in {"выходной", "day off", "off", "вых"}:
        return "Day Off"
    if text in {"на рассмотрении", "pending"}:
        return "Pending"
    if text in {"одобрено", "approved"}:
        return "Approved"
    if text in {"отклонено", "rejected"}:
        return "Rejected"
    return "Working"


def username_of(user) -> str:
    return user.username or ""


def now_stamp() -> str:
    return datetime.now().strftime("%d.%m.%Y %H:%M:%S")


def schedule_menu_keyboard(user_id: int | None = None) -> InlineKeyboardMarkup:
    rows = []
    if user_id is None or not read_user_schedule(user_id):
        rows.append([InlineKeyboardButton(text="📅 Заполнить график", callback_data="sch_submit")])
    rows.extend([
        [InlineKeyboardButton(text="📋 Мой график", callback_data="sch_mine")],
        [InlineKeyboardButton(text="🏖 Запросить выходной", callback_data="sch_dayoff")],
        [InlineKeyboardButton(text="❌ Отмена", callback_data="sch_cancel")],
    ])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def submit_method_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="⌨️ Написать график текстом",
                    callback_data="sch_type",
                )
            ],
            [
                InlineKeyboardButton(
                    text="📆 Заполнить по дням",
                    callback_data="sch_build",
                )
            ],
            [
                InlineKeyboardButton(
                    text="⬅️ Назад",
                    callback_data="sch_menu",
                )
            ],
        ]
    )


def build_days_keyboard(draft: dict) -> InlineKeyboardMarkup:
    rows = []
    row = []

    for day in WEEKDAYS:
        mark = "✅ " if day in draft else ""
        row.append(
            InlineKeyboardButton(
                text=f"{mark}{WEEKDAY_SHORT_RU[day]}",
                callback_data=f"sch_pick_{day}",
            )
        )
        if len(row) == 4:
            rows.append(row)
            row = []

    if row:
        rows.append(row)

    rows.append(
        [
            InlineKeyboardButton(
                text="✅ Подтвердить",
                callback_data="sch_confirm_draft",
            ),
            InlineKeyboardButton(
                text="⬅️ Назад",
                callback_data="sch_submit",
            ),
        ]
    )
    rows.append(
        [
            InlineKeyboardButton(
                text="❌ Отмена",
                callback_data="sch_cancel",
            )
        ]
    )
    return InlineKeyboardMarkup(inline_keyboard=rows)


def confirm_schedule_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="✅ Сохранить график",
                    callback_data="sch_save",
                )
            ],
            [
                InlineKeyboardButton(
                    text="✏️ Изменить",
                    callback_data="sch_submit",
                ),
                InlineKeyboardButton(
                    text="❌ Отмена",
                    callback_data="sch_cancel",
                ),
            ],
        ]
    )


def dayoff_date_keyboard() -> InlineKeyboardMarkup:
    rows = []
    row = []
    today = datetime.now().date()

    for offset in range(14):
        day = today + timedelta(days=offset)
        iso = day.isoformat()
        label = day.strftime("%d.%m")
        if offset == 0:
            label = f"Сегодня {label}"
        row.append(
            InlineKeyboardButton(
                text=label,
                callback_data=f"sch_d_{iso}",
            )
        )
        if len(row) == 3:
            rows.append(row)
            row = []

    if row:
        rows.append(row)

    rows.append(
        [
            InlineKeyboardButton(
                text="⌨️ Ввести дату вручную",
                callback_data="sch_d_manual",
            )
        ]
    )
    rows.append(
        [
            InlineKeyboardButton(
                text="⬅️ Назад",
                callback_data="sch_menu",
            ),
            InlineKeyboardButton(
                text="❌ Отмена",
                callback_data="sch_cancel",
            ),
        ]
    )
    return InlineKeyboardMarkup(inline_keyboard=rows)


def skip_reason_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="⏭ Без причины",
                    callback_data="sch_reason_skip",
                )
            ],
            [
                InlineKeyboardButton(
                    text="❌ Отмена",
                    callback_data="sch_cancel",
                )
            ],
        ]
    )


def confirm_dayoff_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="✅ Отправить заявку",
                    callback_data="sch_do_save",
                )
            ],
            [
                InlineKeyboardButton(
                    text="✏️ Изменить дату",
                    callback_data="sch_dayoff",
                )
            ],
            [
                InlineKeyboardButton(
                    text="❌ Отмена",
                    callback_data="sch_cancel",
                )
            ],
        ]
    )


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def autosize(sheet, widths: dict):
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def ensure_workbook(path: str, title: str, headers: list[str], widths: dict):
    if os.path.exists(path):
        workbook = load_workbook(path)
        sheet = workbook.active
        if sheet.max_row < 1 or sheet.cell(1, 1).value != headers[0]:
            sheet.delete_rows(1, sheet.max_row)
            sheet.append(headers)
            style_header(sheet)
            workbook.save(path)
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    style_header(sheet)
    sheet.freeze_panes = "A2"
    autosize(sheet, widths)
    workbook.save(path)


def ensure_schedule_file():
    ensure_workbook(
        SCHEDULE_FILE,
        "Work Schedule",
        SCHEDULE_HEADERS,
        {
            "A": 18,
            "B": 22,
            "C": 28,
            "D": 16,
            "E": 14,
            "F": 14,
            "G": 14,
            "H": 22,
        },
    )


def ensure_dayoff_file():
    ensure_workbook(
        DAYOFF_FILE,
        "Day Off Requests",
        DAYOFF_HEADERS,
        {
            "A": 14,
            "B": 18,
            "C": 22,
            "D": 28,
            "E": 16,
            "F": 40,
            "G": 14,
            "H": 22,
        },
    )


def format_schedule(draft: dict) -> str:
    lines = []
    for day in WEEKDAYS:
        item = draft.get(day)
        if not item:
            continue
        if item["status"] == "Day Off":
            lines.append(f"{day_ru(day)} — Выходной")
        else:
            lines.append(
                f"{day_ru(day)} — {item['start']}–{item['end']}"
            )
    return "\n".join(lines) if lines else "Дни ещё не указаны."


def parse_time_range(value: str):
    text = value.strip().lower().replace("—", "-").replace("–", "-")
    text = text.replace(".", ":")

    if text in OFF_ALIASES:
        return None, None, "Day Off"

    match = re.fullmatch(
        r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})",
        text,
    )
    if not match:
        return False

    sh, sm, eh, em = (int(part) for part in match.groups())
    if sh > 23 or eh > 23 or sm > 59 or em > 59:
        return False

    start = f"{sh:02d}:{sm:02d}"
    end = f"{eh:02d}:{em:02d}"
    return start, end, "Working"


def parse_schedule_text(text: str) -> dict:
    draft = {}

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        line = line.replace("—", "-").replace("–", "-")
        if "-" not in line:
            raise ValueError(f"Не удалось разобрать строку: {raw_line}")

        day_part, time_part = line.split("-", 1)
        day_key = day_part.strip().lower()
        day = DAY_ALIASES.get(day_key)
        if not day:
            raise ValueError(f"Неизвестный день: {day_part.strip()}")

        parsed = parse_time_range(time_part)
        if parsed is False:
            raise ValueError(f"Неверное время: {time_part.strip()}")

        start, end, status = parsed
        draft[day] = {
            "start": start or "",
            "end": end or "",
            "status": status,
        }

    if not draft:
        raise ValueError("График пустой.")

    return draft


def parse_date(text: str):
    cleaned = text.strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def read_user_schedule(user_id: int) -> dict:
    ensure_schedule_file()
    workbook = load_workbook(SCHEDULE_FILE)
    sheet = workbook.active
    draft = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        if str(row[0]) != str(user_id):
            continue
        day = normalize_day(str(row[3] or ""))
        if day not in WEEKDAYS:
            continue
        draft[day] = {
            "start": str(row[4] or ""),
            "end": str(row[5] or ""),
            "status": normalize_status(str(row[6] or "Working")),
        }

    return draft


async def save_user_schedule(user, draft: dict):
    async with _excel_lock:
        ensure_schedule_file()
        workbook = load_workbook(SCHEDULE_FILE)
        sheet = workbook.active

        rows_to_delete = []
        for idx, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if row and str(row[0]) == str(user.id):
                rows_to_delete.append(idx)

        for idx in reversed(rows_to_delete):
            sheet.delete_rows(idx)

        updated = now_stamp()
        username = username_of(user)
        name = display_name(user)

        for day in WEEKDAYS:
            item = draft.get(day)
            if not item:
                continue
            sheet.append(
                [
                    user.id,
                    username,
                    name,
                    day_ru(day),
                    item.get("start", ""),
                    item.get("end", ""),
                    status_ru(item.get("status", "Working")),
                    updated,
                ]
            )

        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(SCHEDULE_FILE)


def next_request_id(sheet) -> str:
    max_num = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = str(row[0] or "")
        match = re.fullmatch(r"DO(\d+)", value)
        if match:
            max_num = max(max_num, int(match.group(1)))
    return f"DO{max_num + 1:04d}"


async def save_dayoff_request(user, requested_date: str, reason: str) -> str:
    async with _excel_lock:
        ensure_dayoff_file()
        workbook = load_workbook(DAYOFF_FILE)
        sheet = workbook.active
        request_id = next_request_id(sheet)
        sheet.append(
            [
                request_id,
                user.id,
                username_of(user),
                display_name(user),
                requested_date,
                reason,
                status_ru("Pending"),
                now_stamp(),
            ]
        )
        last_row = sheet.max_row
        for cell in sheet[last_row]:
            cell.fill = PENDING_FILL
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(DAYOFF_FILE)
        return request_id


async def update_dayoff_status(request_id: str, status: str) -> Optional[dict]:
    async with _excel_lock:
        ensure_dayoff_file()
        workbook = load_workbook(DAYOFF_FILE)
        sheet = workbook.active
        found = None

        for row in sheet.iter_rows(min_row=2):
            if str(row[0].value or "") != request_id:
                continue
            row[6].value = status_ru(status)
            row[7].value = now_stamp()
            fill = APPROVED_FILL if status == "Approved" else REJECTED_FILL
            for cell in row:
                cell.fill = fill
            found = {
                "request_id": request_id,
                "user_id": row[1].value,
                "date": row[4].value,
                "reason": row[5].value or "",
                "status": status,
            }
            break

        workbook.save(DAYOFF_FILE)
        return found


def user_dayoff_requests(user_id: int) -> list[dict]:
    ensure_dayoff_file()
    workbook = load_workbook(DAYOFF_FILE)
    sheet = workbook.active
    items = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or str(row[1]) != str(user_id):
            continue
        items.append(
            {
                "id": row[0],
                "date": row[4],
                "reason": row[5] or "",
                "status": normalize_status(row[6] or "Pending"),
            }
        )

    return items[-8:]


async def notify_admins(bot, text: str, reply_markup=None):
    for admin_id in _admin_ids:
        try:
            await bot.send_message(
                admin_id,
                text,
                reply_markup=reply_markup,
                parse_mode="HTML",
            )
        except Exception:
            pass


def admin_review_keyboard(request_id: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="✅ Одобрить",
                    callback_data=f"sch_ok_{request_id}",
                ),
                InlineKeyboardButton(
                    text="❌ Отклонить",
                    callback_data=f"sch_no_{request_id}",
                ),
            ]
        ]
    )


async def show_schedule_menu(message: Message, state: FSMContext):
    await state.clear()
    await message.answer(
        "🗓 <b>РАБОЧИЙ ГРАФИК</b>\n\n"
        "Выбери действие:",
        reply_markup=schedule_menu_keyboard(message.from_user.id),
        parse_mode="HTML",
    )


@schedule_router.message(Command("schedule"))
async def schedule_command(message: Message, state: FSMContext):
    if message.chat.type != "private":
        bot_info = await message.bot.get_me()
        await message.reply(
            "Чтобы управлять рабочим графиком, открой бота в личных сообщениях.",
            reply_markup=InlineKeyboardMarkup(
                inline_keyboard=[
                    [
                        InlineKeyboardButton(
                            text="🗓 Открыть график",
                            url=f"https://t.me/{bot_info.username}?start=schedule",
                        )
                    ]
                ]
            ),
        )
        return

    await show_schedule_menu(message, state)


@schedule_router.callback_query(F.data == "sch_menu")
async def schedule_menu_callback(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await show_schedule_menu(callback.message, state)


@schedule_router.callback_query(F.data == "sch_cancel")
async def schedule_cancel_callback(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.answer("Отменено")
    await callback.message.answer("❌ Меню графика закрыто.")


@schedule_router.callback_query(F.data == "sch_submit")
async def schedule_submit_callback(callback: CallbackQuery, state: FSMContext):
    if read_user_schedule(callback.from_user.id):
        await state.clear()
        await callback.answer("График уже заполнен", show_alert=True)
        await callback.message.answer("🔒 <b>График уже заполнен.</b>\n\nСамостоятельно изменить его нельзя. Можно только запросить выходной. Изменения доступны администрации.", reply_markup=schedule_menu_keyboard(callback.from_user.id), parse_mode="HTML")
        return
    if not user_tz_info(callback.from_user.id):
        await state.set_state(ScheduleStates.waiting_city)
        await callback.answer()
        await callback.message.answer("🌍 <b>Сначала укажи свой город</b>\n\nНапиши город, в часовом поясе которого ты живёшь, например: <code>Москва</code>, <code>Киев</code>, <code>Алматы</code>. Бот определит часовой пояс и будет показывать администрации твои часы также по времени Польши.", parse_mode="HTML")
        return
    await state.set_state(None)
    await callback.answer()
    await callback.message.answer(
        "📅 <b>ОТПРАВИТЬ ГРАФИК</b>\n\n"
        "Можно написать всю неделю текстом или заполнить дни по одному.\n\n"
        "<b>Пример:</b>\n"
        "<code>Понедельник — 18:00–23:00\n"
        "Вторник — 17:00–22:00\n"
        "Среда — Выходной\n"
        "Четверг — 18:00–23:00\n"
        "Пятница — 18:00–00:00</code>",
        reply_markup=submit_method_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.message(ScheduleStates.waiting_city, F.text)
async def schedule_city_handler(message: Message, state: FSMContext):
    city=(message.text or "").strip()
    info=await asyncio.to_thread(find_city_timezone, city)
    if not info:
        await message.answer("❌ Не смог найти этот город. Проверь написание и отправь город ещё раз.")
        return
    data=_load_timezones(); data[str(message.from_user.id)]=info; _save_timezones(data)
    await state.set_state(None)
    await message.answer(f"✅ Город: <b>{info['city']}</b> ({info['country']})\n🌍 Часовой пояс: <code>{info['timezone']}</code>\n🇵🇱 Администрации рабочее время будет автоматически показано по времени Польши.\n\nТеперь выбери способ заполнения графика:", reply_markup=submit_method_keyboard(), parse_mode="HTML")

@schedule_router.callback_query(F.data == "sch_type")
async def schedule_type_callback(callback: CallbackQuery, state: FSMContext):
    if read_user_schedule(callback.from_user.id):
        await callback.answer("График уже заполнен", show_alert=True); return
    await state.set_state(ScheduleStates.waiting_schedule_text)
    await callback.answer()
    await callback.message.answer(
        "⌨️ Отправь график текстом.\n\n"
        "Каждый день с новой строки, например:\n"
        "<code>Понедельник — 18:00–23:00\n"
        "Среда — Выходной</code>\n\n"
        "Чтобы отменить, отправь /schedule.",
        parse_mode="HTML",
    )


@schedule_router.message(ScheduleStates.waiting_schedule_text, F.text)
async def schedule_text_handler(message: Message, state: FSMContext):
    if message.text.startswith("/"):
        return

    try:
        draft = parse_schedule_text(message.text)
    except ValueError as error:
        await message.answer(
            f"❌ {error}\n\n"
            "Используй такой формат:\n"
            "<code>Понедельник — 18:00–23:00\n"
            "Вторник — Выходной</code>",
            parse_mode="HTML",
        )
        return

    await state.update_data(draft=draft)
    await state.set_state(None)
    await message.answer(
        "📋 <b>Подтверди график:</b>\n\n"
        f"<code>{format_schedule(draft)}</code>",
        reply_markup=confirm_schedule_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_build")
async def schedule_build_callback(callback: CallbackQuery, state: FSMContext):
    if read_user_schedule(callback.from_user.id):
        await callback.answer("График уже заполнен", show_alert=True); return
    data = await state.get_data()
    draft = data.get("draft") or {}
    await state.update_data(draft=draft)
    await state.set_state(None)
    await callback.answer()
    await callback.message.answer(
        "📆 Нажми на день, затем отправь часы работы или слово <b>Выходной</b>.\n\n"
        f"<code>{format_schedule(draft)}</code>",
        reply_markup=build_days_keyboard(draft),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data.startswith("sch_pick_"))
async def schedule_pick_day(callback: CallbackQuery, state: FSMContext):
    day = callback.data.replace("sch_pick_", "", 1)
    if day not in WEEKDAYS:
        await callback.answer("Неизвестный день", show_alert=True)
        return

    await state.update_data(current_day=day)
    await state.set_state(ScheduleStates.waiting_day_time)
    await callback.answer()
    await callback.message.answer(
        f"🕒 <b>{day_ru(day)}</b>\n\n"
        "Отправь часы работы в формате <code>18:00-23:00</code>\n"
        "или слово <code>Выходной</code>.",
        parse_mode="HTML",
    )


@schedule_router.message(ScheduleStates.waiting_day_time, F.text)
async def schedule_day_time_handler(message: Message, state: FSMContext):
    if message.text.startswith("/"):
        return

    parsed = parse_time_range(message.text)
    if parsed is False:
        await message.answer(
            "❌ Неверное время. Используй <code>18:00-23:00</code> или <code>Выходной</code>.",
            parse_mode="HTML",
        )
        return

    start, end, status = parsed
    data = await state.get_data()
    day = data.get("current_day")
    draft = data.get("draft") or {}

    if day not in WEEKDAYS:
        await state.clear()
        await message.answer("Сессия истекла. Отправь /schedule ещё раз.")
        return

    draft[day] = {
        "start": start or "",
        "end": end or "",
        "status": status,
    }
    await state.update_data(draft=draft, current_day=None)
    await state.set_state(None)
    await message.answer(
        f"✅ {day_ru(day)} сохранён.\n\n"
        f"<code>{format_schedule(draft)}</code>",
        reply_markup=build_days_keyboard(draft),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_confirm_draft")
async def schedule_confirm_draft(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    draft = data.get("draft") or {}

    if not draft:
        await callback.answer("Сначала добавь хотя бы один день.", show_alert=True)
        return

    await callback.answer()
    await callback.message.answer(
        "📋 <b>Подтверди график:</b>\n\n"
        f"<code>{format_schedule(draft)}</code>",
        reply_markup=confirm_schedule_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_save")
async def schedule_save_callback(callback: CallbackQuery, state: FSMContext):
    if read_user_schedule(callback.from_user.id):
        await state.clear(); await callback.answer("График уже был сохранён", show_alert=True); return
    data = await state.get_data()
    draft = data.get("draft") or {}

    if not draft:
        await callback.answer("Нечего сохранять.", show_alert=True)
        return

    try:
        await save_user_schedule(callback.from_user, draft)
        assignments = _load_assignments()
        uid = str(callback.from_user.id)
        if uid not in assignments:
            assignments[uid] = {}
            for day, item in draft.items():
                assignments[uid][day] = {"off": item.get("status") == "Day Off", "start": item.get("start", ""), "end": item.get("end", ""), "duties": []}
            _save_assignments(assignments)
    except Exception as error:
        await callback.answer()
        await callback.message.answer(
            "❌ Не удалось сохранить Excel.\n"
            "Закрой файл, если он открыт, и попробуй снова.\n\n"
            f"<code>{error}</code>",
            parse_mode="HTML",
        )
        return

    await state.clear()
    await callback.answer("Saved")
    await callback.message.answer(
        "✅ <b>График сохранён</b>\n\n"
        f"<code>{format_schedule(draft)}</code>",
        parse_mode="HTML",
    )

    username = username_of(callback.from_user)
    mention = f"@{username}" if username else display_name(callback.from_user)
    await notify_admins(
        callback.bot,
        "📅 <b>График обновлён</b>\n\n"
        f"👤 {mention}\n"
        f"🆔 <code>{callback.from_user.id}</code>\n\n"
        f"<code>{format_schedule(draft)}</code>",
    )


@schedule_router.callback_query(F.data == "sch_mine")
async def my_schedule_callback(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    assigned = _load_assignments().get(str(callback.from_user.id), {})
    draft = read_user_schedule(callback.from_user.id)
    requests = user_dayoff_requests(callback.from_user.id)

    if assigned:
        schedule_text = _week_text(str(callback.from_user.id))
    elif draft:
        schedule_text = format_schedule(draft)
    else:
        schedule_text = "Сохранённого графика пока нет."

    request_lines = []
    for item in requests:
        reason = f" — {item['reason']}" if item["reason"] else ""
        request_lines.append(
            f"{item['id']}: {item['date']} [{status_ru(item['status'])}]{reason}"
        )

    requests_text = (
        "\n".join(request_lines)
        if request_lines
        else "Заявок на выходной нет."
    )

    await callback.message.answer(
        "📋 <b>МОЙ ГРАФИК</b>\n\n"
        f"<code>{schedule_text}</code>\n\n"
        "🏖 <b>Заявки на выходной</b>\n"
        f"<code>{requests_text}</code>",
        reply_markup=schedule_menu_keyboard(callback.from_user.id),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_dayoff")
async def dayoff_start_callback(callback: CallbackQuery, state: FSMContext):
    await state.set_state(None)
    await callback.answer()
    await callback.message.answer(
        "🏖 <b>ЗАПРОС ВЫХОДНОГО</b>\n\n"
        "Выбери дату или введи её вручную.",
        reply_markup=dayoff_date_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_d_manual")
async def dayoff_manual_date(callback: CallbackQuery, state: FSMContext):
    await state.set_state(ScheduleStates.waiting_dayoff_date)
    await callback.answer()
    await callback.message.answer(
        "⌨️ Отправь дату в формате <code>25.09.2026</code> или <code>2026-09-25</code>.",
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data.startswith("sch_d_"))
async def dayoff_picked_date(callback: CallbackQuery, state: FSMContext):
    iso = callback.data.replace("sch_d_", "", 1)
    chosen = parse_date(iso)
    if not chosen:
        await callback.answer("Неверная дата", show_alert=True)
        return

    await state.update_data(dayoff_date=chosen.strftime("%d.%m.%Y"))
    await state.set_state(ScheduleStates.waiting_dayoff_reason)
    await callback.answer()
    await callback.message.answer(
        f"🏖 Дата: <b>{chosen.strftime('%d.%m.%Y')}</b>\n\n"
        "Напиши причину или пропусти этот шаг.",
        reply_markup=skip_reason_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.message(ScheduleStates.waiting_dayoff_date, F.text)
async def dayoff_date_text(message: Message, state: FSMContext):
    if message.text.startswith("/"):
        return

    chosen = parse_date(message.text)
    if not chosen:
        await message.answer(
            "❌ Неверная дата. Используй формат <code>25.09.2026</code>.",
            parse_mode="HTML",
        )
        return

    await state.update_data(dayoff_date=chosen.strftime("%d.%m.%Y"))
    await state.set_state(ScheduleStates.waiting_dayoff_reason)
    await message.answer(
        f"🏖 Дата: <b>{chosen.strftime('%d.%m.%Y')}</b>\n\n"
        "Напиши причину или пропусти этот шаг.",
        reply_markup=skip_reason_keyboard(),
        parse_mode="HTML",
    )


async def ask_dayoff_confirm(target: Message, state: FSMContext, reason: str):
    data = await state.get_data()
    requested = data.get("dayoff_date")
    if not requested:
        await state.clear()
        await target.answer("Сессия истекла. Отправь /schedule ещё раз.")
        return

    await state.update_data(dayoff_reason=reason)
    await state.set_state(None)
    reason_text = reason or "—"
    await target.answer(
        "🏖 <b>Подтверди заявку на выходной</b>\n\n"
        f"📅 Дата: <b>{requested}</b>\n"
        f"💬 Причина: {reason_text}\n"
        f"📌 Статус: <b>{status_ru('Pending')}</b>",
        reply_markup=confirm_dayoff_keyboard(),
        parse_mode="HTML",
    )


@schedule_router.callback_query(F.data == "sch_reason_skip")
async def dayoff_skip_reason(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await ask_dayoff_confirm(callback.message, state, "")


@schedule_router.message(ScheduleStates.waiting_dayoff_reason, F.text)
async def dayoff_reason_text(message: Message, state: FSMContext):
    if message.text.startswith("/"):
        return
    await ask_dayoff_confirm(message, state, message.text.strip())


@schedule_router.callback_query(F.data == "sch_do_save")
async def dayoff_save_callback(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    requested = data.get("dayoff_date")
    reason = data.get("dayoff_reason", "")

    if not requested:
        await callback.answer("Сначала выбери дату.", show_alert=True)
        return

    try:
        request_id = await save_dayoff_request(
            callback.from_user,
            requested,
            reason,
        )
    except Exception as error:
        await callback.answer()
        await callback.message.answer(
            "❌ Не удалось сохранить заявку.\n"
            "Закрой Excel, если файл открыт, и попробуй снова.\n\n"
            f"<code>{error}</code>",
            parse_mode="HTML",
        )
        return

    await state.clear()
    await callback.answer("Отправлено")

    username = username_of(callback.from_user)
    mention = f"@{username}" if username else display_name(callback.from_user)
    reason_text = reason or "—"

    await callback.message.answer(
        "✅ <b>Заявка на выходной отправлена</b>\n\n"
        f"🆔 Номер заявки: <code>{request_id}</code>\n"
        f"📅 Дата: <b>{requested}</b>\n"
        f"💬 Причина: {reason_text}\n"
        f"📌 Статус: <b>{status_ru('Pending')}</b>",
        parse_mode="HTML",
    )

    await notify_admins(
        callback.bot,
        "🏖 <b>Новая заявка на выходной</b>\n\n"
        f"🆔 <code>{request_id}</code>\n"
        f"👤 {mention}\n"
        f"🆔 Пользователь: <code>{callback.from_user.id}</code>\n"
        f"📅 Дата: <b>{requested}</b>\n"
        f"💬 Причина: {reason_text}\n"
        f"📌 Статус: <b>{status_ru('Pending')}</b>",
        reply_markup=admin_review_keyboard(request_id),
    )


@schedule_router.callback_query(F.data.startswith("sch_ok_"))
async def dayoff_approve(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id):
        await callback.answer("Нет доступа", show_alert=True)
        return

    request_id = callback.data.replace("sch_ok_", "", 1)
    found = await update_dayoff_status(request_id, "Approved")
    if not found:
        await callback.answer("Заявка не найдена", show_alert=True)
        return

    await callback.answer("Одобрено")
    await callback.message.answer(
        f"✅ Заявка <code>{request_id}</code> одобрена.",
        parse_mode="HTML",
    )

    try:
        await callback.bot.send_message(
            int(found["user_id"]),
            "✅ <b>Заявка на выходной одобрена</b>\n\n"
            f"🆔 <code>{request_id}</code>\n"
            f"📅 Дата: <b>{found['date']}</b>",
            parse_mode="HTML",
        )
    except Exception:
        pass


@schedule_router.callback_query(F.data.startswith("sch_no_"))
async def dayoff_reject(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id):
        await callback.answer("Нет доступа", show_alert=True)
        return

    request_id = callback.data.replace("sch_no_", "", 1)
    found = await update_dayoff_status(request_id, "Rejected")
    if not found:
        await callback.answer("Заявка не найдена", show_alert=True)
        return

    await callback.answer("Отклонено")
    await callback.message.answer(
        f"❌ Заявка <code>{request_id}</code> отклонена.",
        parse_mode="HTML",
    )

    try:
        await callback.bot.send_message(
            int(found["user_id"]),
            "❌ <b>Заявка на выходной отклонена</b>\n\n"
            f"🆔 <code>{request_id}</code>\n"
            f"📅 Дата: <b>{found['date']}</b>",
            parse_mode="HTML",
        )
    except Exception:
        pass


@schedule_router.callback_query(F.data == "export_schedule_excel")
async def export_schedule_excel(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id):
        await callback.answer("Нет доступа", show_alert=True)
        return

    ensure_schedule_file()
    await callback.answer()
    await callback.message.answer_document(
        document=FSInputFile(SCHEDULE_FILE),
        caption="📅 Рабочие графики",
    )


@schedule_router.callback_query(F.data == "export_dayoff_excel")
async def export_dayoff_excel(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id):
        await callback.answer("Нет доступа", show_alert=True)
        return

    ensure_dayoff_file()
    await callback.answer()
    await callback.message.answer_document(
        document=FSInputFile(DAYOFF_FILE),
        caption="🏖 Заявки на выходной",
    )

# =========================================================
# АДМИН: НАЗНАЧЕНИЕ НЕДЕЛЬНОГО ГРАФИКА И ОБЯЗАННОСТЕЙ
# =========================================================
import json

ASSIGNMENTS_FILE = os.path.join(_BASE_DIR, "admin_assignments.json")
USERS_FILE = os.path.join(_BASE_DIR, "data.json")
DUTIES = {
    "search": "🔎 Искать людей",
    "chat": "🛡 Следить за чатом",
    "tags": "🏷 Выдавать теги",
    "owner_tech": "👑 Владелец / 🛠 Тех админ",
}

class AdminScheduleStates(StatesGroup):
    waiting_hours = State()


def _load_assignments() -> dict:
    try:
        with open(ASSIGNMENTS_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _save_assignments(data: dict):
    with open(ASSIGNMENTS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _known_users() -> list[dict]:
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            users = json.load(f).get("users", {})
    except (OSError, json.JSONDecodeError):
        users = {}
    result = []
    for uid, item in users.items():
        result.append({"id": str(item.get("id") or uid), "name": item.get("name") or item.get("username") or uid,
                       "username": item.get("username") or ""})
    return result


def _day_record(data: dict, uid: str, day: str) -> dict:
    return data.setdefault(uid, {}).setdefault(day, {"off": False, "start": "", "end": "", "duties": []})


def admin_schedule_users_keyboard() -> InlineKeyboardMarkup:
    rows = []
    for u in _known_users():
        label = u["name"][:28]
        if u["username"]:
            label += f" (@{u['username']})"
        rows.append([InlineKeyboardButton(text=label[:55], callback_data=f"adm_su_{u['id']}")])
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin_panel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def admin_schedule_days_keyboard(uid: str) -> InlineKeyboardMarkup:
    data = _load_assignments()
    rows, row = [], []
    for day in WEEKDAYS:
        rec = data.get(uid, {}).get(day, {})
        mark = "🏖" if rec.get("off") else ("✅" if rec.get("duties") or rec.get("start") else "▫️")
        row.append(InlineKeyboardButton(text=f"{mark} {WEEKDAY_SHORT_RU[day]}", callback_data=f"adm_sd_{uid}_{day[:3]}"))
        if len(row) == 2:
            rows.append(row); row = []
    if row: rows.append(row)
    rows.append([InlineKeyboardButton(text="♾ 24/7 · Владелец/Тех админ", callback_data=f"adm_247_{uid}")])
    rows.append([InlineKeyboardButton(text="📋 Показать неделю", callback_data=f"adm_sw_{uid}")])
    rows.append([InlineKeyboardButton(text="🗑 Удалить весь график", callback_data=f"adm_del_{uid}")])
    rows.append([InlineKeyboardButton(text="⬅️ К сотрудникам", callback_data="admin_schedule")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _day_from_short(short: str) -> str:
    return next((d for d in WEEKDAYS if d.startswith(short)), "")


def admin_day_keyboard(uid: str, day: str) -> InlineKeyboardMarkup:
    data = _load_assignments(); rec = data.get(uid, {}).get(day, {})
    duties = rec.get("duties", [])
    rows = [
        [InlineKeyboardButton(text=("☑️ " if "search" in duties else "⬜ ") + DUTIES["search"], callback_data=f"adm_dt_{uid}_{day[:3]}_search")],
        [InlineKeyboardButton(text=("☑️ " if "chat" in duties else "⬜ ") + DUTIES["chat"], callback_data=f"adm_dt_{uid}_{day[:3]}_chat")],
        [InlineKeyboardButton(text=("☑️ " if "tags" in duties else "⬜ ") + DUTIES["tags"], callback_data=f"adm_dt_{uid}_{day[:3]}_tags")],
        [InlineKeyboardButton(text=("☑️ " if "owner_tech" in duties else "⬜ ") + DUTIES["owner_tech"], callback_data=f"adm_dt_{uid}_{day[:3]}_owner_tech")],
        [InlineKeyboardButton(text="⏰ Изменить рабочие часы", callback_data=f"adm_hr_{uid}_{day[:3]}")],
        [InlineKeyboardButton(text="🏖 Поставить/убрать выходной", callback_data=f"adm_off_{uid}_{day[:3]}")],
        [InlineKeyboardButton(text="🗑 Очистить этот день", callback_data=f"adm_cl_{uid}_{day[:3]}")],
        [InlineKeyboardButton(text="⬅️ К дням недели", callback_data=f"adm_su_{uid}")],
    ]
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _user_label(uid: str) -> str:
    for u in _known_users():
        if u["id"] == str(uid): return u["name"]
    return uid


def _day_text(uid: str, day: str) -> str:
    rec = _load_assignments().get(str(uid), {}).get(day, {})
    if rec.get("always"):
        info = "♾ <b>Работает 24/7</b>\n📌 Обязанности: 👑 Владелец / 🛠 Тех админ"
    elif rec.get("off"):
        info = "🏖 <b>Выходной</b>"
    else:
        hours = f"{rec.get('start')}–{rec.get('end')}" if rec.get("start") and rec.get("end") else "не указаны"
        duties = rec.get("duties", [])
        duty_text = ", ".join(DUTIES.get(x, x) for x in duties) if duties else "не назначены"
        info = f"⏰ Часы: <b>{hours}</b>\n📌 Обязанности: {duty_text}"
    return f"👤 <b>{_user_label(uid)}</b>\n📅 <b>{day_ru(day)}</b>\n\n{info}"


def _week_text(uid: str) -> str:
    data = _load_assignments().get(str(uid), {})
    lines = [f"📋 <b>График: {_user_label(uid)}</b>", ""]
    for day in WEEKDAYS:
        rec = data.get(day, {})
        if rec.get("always"):
            val = "♾ 24/7 · Владелец / Тех админ"
        elif rec.get("off"):
            val = "🏖 Выходной"
        else:
            hours = f"{rec.get('start')}–{rec.get('end')}" if rec.get("start") and rec.get("end") else "часы не указаны"
            duties = ", ".join(DUTIES.get(x, x).split(" ",1)[-1] for x in rec.get("duties", [])) or "без обязанностей"
            val = f"{hours} · {duties}"
        lines.append(f"<b>{WEEKDAY_SHORT_RU[day]}</b>: {val}")
    return "\n".join(lines)


@schedule_router.callback_query(F.data == "admin_schedule")
async def admin_schedule_panel(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id):
        await callback.answer("Нет доступа", show_alert=True); return
    await callback.answer()
    await callback.message.answer("🗓 <b>УПРАВЛЕНИЕ ГРАФИКОМ</b>\n\nВыбери сотрудника:", reply_markup=admin_schedule_users_keyboard(), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_su_"))
async def admin_select_user(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    uid = callback.data[7:]
    await callback.answer()
    await callback.message.answer(f"👤 <b>{_user_label(uid)}</b>\n\nВыбери день недели:", reply_markup=admin_schedule_days_keyboard(uid), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_sd_"))
async def admin_select_day(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    _, _, uid, short = callback.data.split("_", 3); day = _day_from_short(short)
    await callback.answer(); await callback.message.answer(_day_text(uid, day), reply_markup=admin_day_keyboard(uid, day), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_dt_"))
async def admin_toggle_duty(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    _, _, uid, short, duty = callback.data.split("_", 4); day = _day_from_short(short)
    data = _load_assignments(); rec = _day_record(data, uid, day); rec["off"] = False; rec["always"] = False
    if duty in rec["duties"]: rec["duties"].remove(duty)
    else: rec["duties"].append(duty)
    _save_assignments(data); await callback.answer("Сохранено")
    await callback.message.edit_text(_day_text(uid, day), reply_markup=admin_day_keyboard(uid, day), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_off_"))
async def admin_toggle_off(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    _, _, uid, short = callback.data.split("_", 3); day = _day_from_short(short)
    data = _load_assignments(); rec = _day_record(data, uid, day); rec["always"] = False; rec["off"] = not rec.get("off", False)
    _save_assignments(data); await callback.answer("Выходной изменён")
    await callback.message.edit_text(_day_text(uid, day), reply_markup=admin_day_keyboard(uid, day), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_cl_"))
async def admin_clear_day(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    _, _, uid, short = callback.data.split("_", 3); day = _day_from_short(short)
    data = _load_assignments(); data.setdefault(uid, {})[day] = {"off": False, "start": "", "end": "", "duties": []}; _save_assignments(data)
    await callback.answer("День очищен")
    await callback.message.edit_text(_day_text(uid, day), reply_markup=admin_day_keyboard(uid, day), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_hr_"))
async def admin_set_hours_start(callback: CallbackQuery, state: FSMContext):
    if not is_schedule_admin(callback.from_user.id): return
    _, _, uid, short = callback.data.split("_", 3); day = _day_from_short(short)
    await state.set_state(AdminScheduleStates.waiting_hours); await state.update_data(admin_uid=uid, admin_day=day)
    await callback.answer(); await callback.message.answer("⏰ Отправь рабочие часы в формате <code>18:00-23:00</code>\nДля удаления часов отправь <code>-</code>.", parse_mode="HTML")


@schedule_router.message(AdminScheduleStates.waiting_hours)
async def admin_set_hours_finish(message: Message, state: FSMContext):
    if not is_schedule_admin(message.from_user.id): await state.clear(); return
    ctx = await state.get_data(); uid, day = ctx["admin_uid"], ctx["admin_day"]
    text = (message.text or "").strip()
    if text == "-": start = end = ""
    else:
        m = re.fullmatch(r"\s*([01]?\d|2[0-3]):([0-5]\d)\s*[-–—]\s*([01]?\d|2[0-3]):([0-5]\d)\s*", text)
        if not m:
            await message.answer("❌ Неверный формат. Пример: <code>18:00-23:00</code>", parse_mode="HTML"); return
        start = f"{int(m.group(1)):02d}:{m.group(2)}"; end = f"{int(m.group(3)):02d}:{m.group(4)}"
    data = _load_assignments(); rec = _day_record(data, uid, day); rec.update({"start": start, "end": end, "off": False, "always": False}); _save_assignments(data)
    await state.clear(); await message.answer("✅ Рабочие часы сохранены.\n\n" + _day_text(uid, day), reply_markup=admin_day_keyboard(uid, day), parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_247_"))
async def admin_set_247(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    uid = callback.data[8:]
    data = _load_assignments()
    for day in WEEKDAYS:
        data.setdefault(str(uid), {})[day] = {
            "off": False, "start": "", "end": "",
            "duties": ["owner_tech"], "always": True
        }
    _save_assignments(data)
    await callback.answer("Установлено: 24/7")
    await callback.message.answer(
        f"♾ <b>{_user_label(uid)}</b> назначен на режим <b>24/7</b>.\n"
        "📌 Обязанности: Владелец / Тех админ",
        reply_markup=admin_schedule_days_keyboard(uid), parse_mode="HTML"
    )


@schedule_router.callback_query(F.data.startswith("adm_sw_"))
async def admin_show_week(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    uid = callback.data[7:]; await callback.answer()
    assigned=_load_assignments().get(str(uid), {})
    if assigned:
        text=_week_text(uid)
    else:
        draft=read_user_schedule(int(uid))
        info=user_tz_info(uid)
        text=f"📋 <b>График: {_user_label(uid)}</b>\n"
        if info: text += f"🌍 {info.get('city','')} · <code>{info.get('timezone','')}</code>\n🇵🇱 Время ниже переведено на польское:\n\n"
        text += f"<code>{format_schedule_poland(draft, uid)}</code>" if draft else "График не заполнен."
    await callback.message.answer(text, parse_mode="HTML")


@schedule_router.callback_query(F.data.startswith("adm_del_"))
async def admin_delete_schedule_confirm(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    uid=callback.data[8:]; await callback.answer()
    kb=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="✅ Да, удалить",callback_data=f"adm_delok_{uid}"),InlineKeyboardButton(text="❌ Отмена",callback_data=f"adm_su_{uid}")]])
    await callback.message.answer(f"⚠️ Удалить весь график <b>{_user_label(uid)}</b>? После удаления человек сможет заполнить его заново один раз.",reply_markup=kb,parse_mode="HTML")

@schedule_router.callback_query(F.data.startswith("adm_delok_"))
async def admin_delete_schedule(callback: CallbackQuery):
    if not is_schedule_admin(callback.from_user.id): return
    uid=callback.data[10:]
    async with _excel_lock:
        ensure_schedule_file(); wb=load_workbook(SCHEDULE_FILE); sh=wb.active
        for i in range(sh.max_row,1,-1):
            if str(sh.cell(i,1).value)==str(uid): sh.delete_rows(i)
        wb.save(SCHEDULE_FILE)
    data=_load_assignments(); data.pop(str(uid),None); _save_assignments(data)
    tz=_load_timezones(); tz.pop(str(uid),None); _save_timezones(tz)
    await callback.answer("График удалён")
    await callback.message.answer(f"🗑 График <b>{_user_label(uid)}</b> полностью удалён. Пользователь может заполнить его заново.",parse_mode="HTML")


def personal_schedule_choice_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💬 Показать сообщением", callback_data="sch_personal_text")],
        [InlineKeyboardButton(text="📊 Скачать мой график Excel", callback_data="sch_personal_excel")],
    ])


def _personal_schedule_rows(uid: str) -> list[dict]:
    assigned = _load_assignments().get(str(uid), {})
    rows = []
    if assigned:
        for day in WEEKDAYS:
            rec = assigned.get(day, {})
            if rec.get("always"):
                rows.append({"day": day, "hours": "24/7", "status": "Работа 24/7", "duties": "Владелец / Тех админ"})
            elif rec.get("off"):
                rows.append({"day": day, "hours": "—", "status": "Выходной", "duties": "—"})
            else:
                hours = f"{rec.get('start')}–{rec.get('end')}" if rec.get("start") and rec.get("end") else "Не указаны"
                duties = ", ".join(DUTIES.get(x, x).split(" ", 1)[-1] for x in rec.get("duties", [])) or "Без обязанностей"
                rows.append({"day": day, "hours": hours, "status": "Рабочий день", "duties": duties})
        return rows

    submitted = read_user_schedule(int(uid))
    for day in WEEKDAYS:
        rec = submitted.get(day, {})
        if not rec:
            rows.append({"day": day, "hours": "—", "status": "Не заполнено", "duties": "—"})
        elif rec.get("status") == "Day Off":
            rows.append({"day": day, "hours": "—", "status": "Выходной", "duties": "—"})
        else:
            hours = f"{rec.get('start')}–{rec.get('end')}" if rec.get("start") and rec.get("end") else "Не указаны"
            rows.append({"day": day, "hours": hours, "status": "Рабочий день", "duties": "Не назначены"})
    return rows


def create_personal_schedule_excel(user_id: int) -> str:
    uid = str(user_id)
    rows = _personal_schedule_rows(uid)
    path = os.path.join(_BASE_DIR, f"shinobi_schedule_{uid}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Мой график"
    ws.merge_cells("A1:D1")
    ws["A1"] = f"SHINOBI TEAM — График: {_user_label(uid)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["День недели", "Рабочее время", "Статус", "Обязанности"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(3, col, value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for r_idx, item in enumerate(rows, 4):
        values = [WEEKDAY_RU[item["day"]], item["hours"], item["status"], item["duties"]]
        for c_idx, value in enumerate(values, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {"A": 18, "B": 20, "C": 18, "D": 42}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A4"
    wb.save(path)
    return path


async def send_personal_schedule_text(message: Message, user_id: int):
    uid = str(user_id)
    assigned = _load_assignments().get(uid, {})
    submitted = read_user_schedule(user_id)
    if assigned:
        await message.answer(_week_text(uid), parse_mode="HTML")
    elif submitted:
        await message.answer("📋 <b>МОЙ ГРАФИК</b>\n\n<code>" + format_schedule(submitted) + "</code>", parse_mode="HTML")
    else:
        await message.answer("📋 Твой график ещё не заполнен. Открой бота в ЛС и заполни его один раз.")


@schedule_router.callback_query(F.data == "sch_personal_text")
async def personal_schedule_text_callback(callback: CallbackQuery):
    await callback.answer()
    await send_personal_schedule_text(callback.message, callback.from_user.id)


@schedule_router.callback_query(F.data == "sch_personal_excel")
async def personal_schedule_excel_callback(callback: CallbackQuery):
    uid = callback.from_user.id
    if not _load_assignments().get(str(uid), {}) and not read_user_schedule(uid):
        await callback.answer("График ещё не заполнен", show_alert=True)
        return
    await callback.answer("Формирую Excel…")
    path = create_personal_schedule_excel(uid)
    try:
        await callback.message.answer_document(
            document=FSInputFile(path),
            caption="📊 <b>Твой персональный график Shinobi Team</b>\n\nВ файле находится только твой график.",
            parse_mode="HTML",
        )
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


@schedule_router.message(F.text.casefold().in_({".мой график", "мой график"}))
async def personal_week_schedule(message: Message):
    await message.answer(
        "📋 <b>МОЙ ГРАФИК</b>\n\nВыбери, как получить свой персональный график:",
        reply_markup=personal_schedule_choice_keyboard(),
        parse_mode="HTML",
    )


def all_schedules_choice_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="💬 Показать сообщением", callback_data="sch_all_text")],
        [InlineKeyboardButton(text="📊 Скачать общий график Excel", callback_data="sch_all_excel")],
    ])


def _all_schedules_text() -> str:
    users = _known_users()
    if not users:
        return "🗓 <b>ОБЩИЙ ГРАФИК SHINOBI TEAM</b>\n\nПользователей пока нет."
    blocks = ["🗓 <b>ОБЩИЙ ГРАФИК SHINOBI TEAM</b>"]
    assignments = _load_assignments()
    for u in users:
        uid = str(u["id"])
        assigned = assignments.get(uid, {})
        submitted = read_user_schedule(int(uid))
        if not assigned and not submitted:
            continue
        label = u["name"] + (f" (@{u['username']})" if u.get("username") else "")
        if assigned:
            body = _week_text(uid)
            # _week_text уже содержит заголовок пользователя — оставляем как отдельный блок.
            blocks.append(body)
        else:
            blocks.append(f"👤 <b>{label}</b>\n<code>{format_schedule_poland(submitted, uid)}</code>")
    if len(blocks) == 1:
        blocks.append("Ни один пользователь ещё не заполнил график.")
    return "\n\n━━━━━━━━━━━━━━\n\n".join(blocks)


@schedule_router.callback_query(F.data == "sch_all_text")
async def all_schedules_text_callback(callback: CallbackQuery):
    await callback.answer()
    text = _all_schedules_text()
    # Telegram ограничивает сообщение примерно 4096 символами.
    if len(text) <= 3900:
        await callback.message.answer(text, parse_mode="HTML")
    else:
        # Делим по блокам, сохраняя HTML-разметку каждого блока.
        for part in text.split("\n\n━━━━━━━━━━━━━━\n\n"):
            await callback.message.answer(part[:3900], parse_mode="HTML")


@schedule_router.callback_query(F.data == "sch_all_excel")
async def all_schedules_excel_callback(callback: CallbackQuery):
    await callback.answer("Формирую общий Excel…")
    ensure_schedule_file()
    await callback.message.answer_document(FSInputFile(SCHEDULE_FILE), caption="📅 <b>Общий график Shinobi Team</b>\nГрафики всех пользователей бота.", parse_mode="HTML")


@schedule_router.message(F.text.in_({".график", "график", "График", ".График"}))
async def public_week_schedule(message: Message):
    await message.answer(
        "🗓 <b>SHINOBI TEAM · ОБЩИЙ ГРАФИК</b>\n\nВыбери, как показать график всех пользователей бота:",
        reply_markup=all_schedules_choice_keyboard(),
        parse_mode="HTML",
    )
